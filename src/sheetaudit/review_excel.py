"""Human-editable workbook interchange, with explicit deletions and version binding."""
from datetime import date, datetime
from decimal import Decimal
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .i18n import translate as tr, current_language, check_language
from .model import FIELDS, AMOUNTS, NUMBERS, canonical_number, identified_issues, json_hash
from .report import sheet, append, finish
from .review import review_template

ITEM_HEADERS = ["行 ID", "页码", "品名", "数量", "单价", "金额", "原文证据", "操作"]
FIELD_HEADERS = ["字段", "含义", "原值", "修正值（空白=未知）"]
LABELS = dict(zip(FIELDS, ["单据类型", "单据号", "日期", "销售方", "购买方", "币种", "小计", "税额", "折扣", "运费", "总计"]))


def write_review_workbook(document, path, source=None):
    wb = Workbook()
    wb.remove(wb.active)
    sign = sheet(wb, tr("复核签注"), [tr("项目"), tr("填写值"), tr("说明")], [27, 65, 100],
                 tr("先在本机打开原 PDF，再编辑黄色单元格。保存后执行 review 命令；Excel 中不自动重算。"))
    rows = [
        ["reviewer", "", tr("填写实际复核人，不由代理代签")],
        ["reviewed_at", "", tr("带时区时间，例如 2026-10-25T10:00:00+08:00")],
        ["coverage_confirmed", tr("未确认"), tr("逐页确认行数、字段和税费口径后选择“已确认”")],
        ["notes", "", tr("本次核对与修改的依据，必填")],
        ["source_file", document["source"]["file"], tr("原 PDF 文件名；下面命令的 --source 指向该原件")],
        ["rounding", "ROUND_HALF_UP", tr("金额采用指定小数位四舍五入；不自动推断业务口径")],
        ["decimal_places", document.get("audit_policy", {}).get("decimal_places", 2), tr("生成时小数位；默认在重算时沿用")],
        ["tolerance", document.get("audit_policy", {}).get("tolerance", "0.01"), tr("生成时容差；需要改变时在 review 命令显式指定")],
    ]
    for row in rows:
        append(sign, row)
    confirmation = DataValidation(type="list", formula1=tr('"未确认,已确认"'))
    confirmation.errorTitle = tr("请选择确认状态")
    confirmation.error = tr("仅支持未确认或已确认")
    confirmation.showErrorMessage = True
    sign.add_data_validation(confirmation)
    confirmation.add(sign["B6"])
    if source:
        sign["B8"].hyperlink = source.resolve().as_uri()

    fields = sheet(wb, tr("字段修正"), FIELD_HEADERS, [24, 18, 50, 50],
                   tr("仅编辑 D 列；未知留空，确认无此项时填 0。标识符和金额按文本保留，避免丢失前导零。"))
    for key in FIELDS:
        append(fields, [key, tr(LABELS[key]), document["fields"][key], document["fields"][key]])
    items = sheet(wb, tr("明细修正"), ITEM_HEADERS, [18, 10, 36, 18, 18, 18, 85, 12],
                  tr("编辑现有行时保留行 ID；删行请选择“删除”，不要直接移除整行。新增行须填写新的唯一 ID、页码和证据。"))
    originals = sheet(wb, tr("原始明细"), ITEM_HEADERS[:-1], [18, 10, 36, 18, 18, 18, 85],
                      tr("本次修正前的数值快照，不可编辑；更早版本保存在 before-review.json。"))
    for row in document["items"]:
        values = [row[k] for k in ("row_id", "page", "description", *NUMBERS, "evidence")]
        append(items, [*values, tr("保留")])
        append(originals, values)
    action = DataValidation(type="list", formula1=tr('"保留,删除"'))
    action.showErrorMessage = True
    items.add_data_validation(action)
    action.add("H4:H10003")
    notices = sheet(wb, tr("问题处理"), [tr("问题 ID"), tr("检查代码"), tr("页码"), tr("行 ID"), tr("提示"), tr("处理状态"), tr("处理依据")],
                    [28, 26, 9, 16, 70, 16, 70],
                    tr("提取提示可标记已解决并填写依据；算术差异必须通过数据修正后重新计算，不允许手动隐藏。"))
    for entry in identified_issues(document):
        append(notices, [entry["issue_id"], entry["code"], entry["page"], entry["row_id"], tr(entry["message"]), tr("待处理"), ""])
    resolution = DataValidation(type="list", formula1=tr('"待处理,已解决"'))
    resolution.showErrorMessage = True
    notices.add_data_validation(resolution)
    resolution.add(f"F4:F{max(4, notices.max_row)}")
    for ws, editable in [(sign, (2,)), (fields, (4,)), (items, tuple(range(1, 9))), (originals, ()), (notices, (6, 7))]:
        finish(ws, review_cols=editable)
        ws.freeze_panes = "C4" if ws is not sign else "B4"
        # Protection prevents casual edits only; hashes and reader validation enforce binding.
        ws.protection.sheet = ws is not items
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if cell.column in editable and not (ws is sign and cell.row > 7):
                    cell.protection = Protection(locked=False)
                    cell.number_format = "@" if not (ws is items and cell.column == 2) else "0"
    meta = wb.create_sheet("_meta")
    for key, value in {"format": "sheetaudit-review-v1", "language": current_language(), "source_sha256": document["source"]["sha256"],
                       "base_extraction_sha256": json_hash(document)}.items():
        append(meta, [key, value])
    meta.sheet_state = "veryHidden"
    wb.save(path)


def read_review_workbook(path, document):
    if path.stat().st_size > 20_000_000:
        raise ValueError("复核工作簿超过 20 MB")
    with ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 60_000_000:
            raise ValueError("复核工作簿解压体积过大")
    wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    try:
        if "_meta" not in wb.sheetnames or wb["_meta"].max_row > 10 or wb["_meta"].max_column != 2:
            raise ValueError("复核工作簿元信息无效")
        meta = dict(wb["_meta"].iter_rows(values_only=True))
        language = check_language(meta.get("language", "zh"))
        def name(value):
            return tr(value, language)
        required = {"复核签注", "字段修正", "明细修正", "原始明细", "问题处理", "_meta"}
        if set(wb.sheetnames) != {name(value) for value in required}:
            raise ValueError("复核工作簿缺少工作表或包含未知工作表")
        headers = {"复核签注": ["项目", "填写值", "说明"], "字段修正": FIELD_HEADERS,
                   "明细修正": ITEM_HEADERS, "原始明细": ITEM_HEADERS[:-1],
                   "问题处理": ["问题 ID", "检查代码", "页码", "行 ID", "提示", "处理状态", "处理依据"]}
        for sheet_name, expected_headers in headers.items():
            actual_headers = list(next(wb[name(sheet_name)].iter_rows(min_row=3, max_row=3, values_only=True)))
            if actual_headers != [name(value) for value in expected_headers]:
                raise ValueError("复核工作簿列顺序或表头被修改")
        for ws in wb:
            if ws.max_row > 10010 or ws.max_column > 16:
                raise ValueError("复核工作簿行列超出限制")
            for row in ws:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("复核工作簿不接受公式；请粘贴明确数值")
        expected = {"format": "sheetaudit-review-v1", "source_sha256": document["source"]["sha256"],
                    "base_extraction_sha256": json_hash(document)}
        if "language" in meta:
            expected["language"] = language
        if meta != expected:
            raise ValueError("复核工作簿与当前原提取版本不匹配")
        out = review_template(document)
        sign = {row[0]: row[1] for row in wb[name("复核签注")].iter_rows(min_row=4, values_only=True)}
        for key in ("reviewer", "reviewed_at", "notes"):
            out[key] = sign.get(key) or ""
        coverage = sign.get("coverage_confirmed")
        if coverage not in (name("已确认"), name("未确认")):
            raise ValueError("复核完整性请选择已确认或未确认")
        out["coverage_confirmed"] = coverage == name("已确认")
        rows = list(wb[name("字段修正")].iter_rows(min_row=4, values_only=True))
        if len(rows) != len(FIELDS) or {r[0] for r in rows} != set(FIELDS):
            raise ValueError("字段行被删除、重复或字段名被更改")
        for key, label, original_value, value in rows:
            if original_value != document["fields"][key]:
                raise ValueError("原值列被修改；请仅编辑字段修正列")
            if key in AMOUNTS:
                value = excel_number(value)
            elif key == "date" and isinstance(value, (date, datetime)):
                value = value.strftime("%Y-%m-%d")
            elif value is not None and not isinstance(value, str):
                raise ValueError("单据号、名称等字段必须是文本；请避免 Excel 自动转换标识符")
            out["fields"][key] = value if value != "" else None
        original_rows = [[row[k] for k in ("row_id", "page", "description", *NUMBERS, "evidence")]
                         for row in document["items"]]
        if [list(row) for row in wb[name("原始明细")].iter_rows(min_row=4, values_only=True)] != original_rows:
            raise ValueError("原始明细被修改；请只编辑明细修正工作表")
        items, seen = [], set()
        for row in wb[name("明细修正")].iter_rows(min_row=4, max_col=8, values_only=True):
            if all(value is None for value in row):
                continue
            identity, page, desc, q, p, a, evidence, action = row
            if not isinstance(identity, str) or identity in seen:
                raise ValueError("行 ID 缺失或重复")
            seen.add(identity)
            if action not in (name("保留"), name("删除")):
                raise ValueError("每条明细的操作必须为保留或删除")
            if action == name("删除"):
                continue
            if isinstance(page, str) and page.isdigit():
                page = int(page)
            items.append({"row_id": identity, "page": page, "description": desc,
                          "quantity": excel_number(q), "unit_price": excel_number(p),
                          "amount": excel_number(a), "evidence": evidence})
        if not {row["row_id"] for row in document["items"]} <= seen:
            raise ValueError("原明细行被直接删除或 ID 被改动；请使用操作列的删除选项")
        out["items"] = items
        known = {entry["issue_id"]: entry for entry in identified_issues(document)}
        out["resolutions"] = []
        seen = set()
        for identity, code, page, row_id, message, state, note in wb[name("问题处理")].iter_rows(min_row=4, max_col=7, values_only=True):
            if identity not in known or identity in seen:
                raise ValueError("问题 ID 被修改或重复")
            seen.add(identity)
            original = known[identity]
            if [code, page, row_id, message] != [original["code"], original["page"], original["row_id"], name(original["message"])]:
                raise ValueError("原问题内容被修改；请只编辑处理状态和依据")
            if state not in (name("待处理"), name("已解决")):
                raise ValueError("问题状态必须为待处理或已解决")
            out["resolutions"].append({"issue_id": identity, "action": "resolved" if state == name("已解决") else "open", "note": note or ""})
        if seen != set(known):
            raise ValueError("原问题行被删除")
        return out
    finally:
        wb.close()


def excel_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError("金额或数量必须为明确的十进制数值")
    if isinstance(value, str):
        return canonical_number(value.strip())
    return canonical_number(format(Decimal(str(value)), "f"))
