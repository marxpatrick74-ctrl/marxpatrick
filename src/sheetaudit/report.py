"""Portable XLSX snapshots; untrusted document text is always stored as text."""
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from .i18n import translate as tr
from .model import AMOUNTS

NAVY = "183349"
TEAL = "007F82"
RED = "FCE4E4"
YELLOW = "FFF2CC"
PALE = "EEF4F8"


def append(ws, values):
    ws.append([ILLEGAL_CHARACTERS_RE.sub("", v)[:32767] if isinstance(v, str) else v for v in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # Disarm =, +, -, @ and external-reference formulas.


def num(value):
    return Decimal(value) if value is not None else None


def sheet(wb, name, headers, widths, subtitle):
    name, headers, subtitle = tr(name), [tr(h) for h in headers], tr(subtitle)
    ws = wb.create_sheet(name)
    append(ws, ["SheetAudit AI · " + name])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    append(ws, [subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    append(ws, headers)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 32
    for row in (1, 3):
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=NAVY if row == 1 else TEAL)
            cell.font = Font(name="Microsoft YaHei", size=13 if row == 1 else 10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws["A2"].font = Font(name="Microsoft YaHei", size=10, color="526879")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D4"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    return ws


def finish(ws, number_cols=(), review_cols=(), integer_cols=(), general_cols=()):
    for row in ws.iter_rows(min_row=4):
        ws.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="183349")
            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            if cell.column in number_cols:
                cell.number_format = '#,##0.00######;[Red](#,##0.00######);0.00'
            if cell.column in integer_cols:
                cell.number_format = '#,##0'
            if cell.column in general_cols:
                cell.number_format = 'General'
            if cell.column in review_cols:
                cell.fill = PatternFill("solid", fgColor=YELLOW)
            if cell.value == tr("差异"):
                cell.fill = PatternFill("solid", fgColor=RED)
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="A32323")
            elif cell.value == tr("无法核验"):
                cell.fill = PatternFill("solid", fgColor=YELLOW)
    ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{max(3, ws.max_row)}"
    ws.print_options.horizontalCentered = True


def write_report(records, errors, output):
    wb = Workbook()
    wb.remove(wb.active)
    overview = sheet(wb, tr("单据汇总"),
        [tr("单据 ID"), tr("文件"), tr("单据号"), tr("类型"), tr("币种"), tr("明细数"), tr("差异数"), tr("无法核验"), tr("提取提示数"), tr("人工复核"), tr("处理状态")],
        [19, 30, 23, 20, 10, 10, 10, 12, 12, 34, 22],
        tr("固定快照，不跨币种汇总金额。算术一致不代表发票真实、提取完整或可入账。修正数据须通过 review 命令重跑。"))
    detail = sheet(wb, tr("明细核对"),
        [tr("单据 ID"), tr("行 ID"), tr("页码"), tr("品名"), tr("数量"), tr("单价"), tr("原金额"), tr("计算金额"), tr("差异"), tr("结果"), tr("币种"), tr("原文证据")],
        [19, 15, 8, 28, 13, 15, 15, 15, 14, 13, 9, 65],
        tr("差异=原金额−四舍五入后的数量×单价；计算由 Python Decimal 完成。表内数字为运行时快照，改 Excel 不会重算。"))
    queue = sheet(wb, tr("差异与复核"),
        [tr("单据 ID"), tr("行 ID"), tr("页码"), tr("检查代码"), tr("检查/原因"), tr("计算值"), tr("原值"), tr("差异"), tr("结果"), tr("说明"), tr("复核人"), tr("处理备注")],
        [19, 15, 8, 25, 37, 15, 15, 14, 14, 64, 17, 36],
        tr("此表用于查看结果；实际修正请编辑对应单据的 review.xlsx。已解决提示保留处理依据，计算差异不能手动关闭。"))
    fields = sheet(wb, tr("字段与来源"),
        [tr("单据 ID"), tr("字段"), tr("提取值"), tr("来源页码"), tr("来源证据/元信息")],
        [19, 24, 48, 16, 100],
        tr("source.json 保存原始页面文字与表格，extracted.json 保存提取结果。SHA-256 用于绑定人工修正与原 PDF。"))
    for record in records:
        d, a = record["document"], record["audit"]
        did = d["document_id"]
        f = d["fields"]
        append(overview, [did, d["source"]["file"], f["document_number"], f["document_type"], f["currency"],
                          len(d["items"]), a["differences"], a["uncheckable"], len(a["issues"]), tr(a["review_status"]), tr("报告已生成")])
        line_checks = {c["row_id"]: c for c in a["checks"] if c["code"] == "LINE_AMOUNT"}
        for item in d["items"]:
            c = line_checks[item["row_id"]]
            append(detail, [did, item["row_id"], item["page"], item["description"], num(item["quantity"]),
                            num(item["unit_price"]), num(item["amount"]), num(c["expected"]), num(c["delta"]),
                            tr(c["status"]), f["currency"], item["evidence"]])
        for c in a["checks"]:
            # Keep all document-level totals so skipped checks cannot disappear.
            if c["status"] != "算术一致" or c["row_id"] is None:
                append(queue, [did, c["row_id"], c["page"], c["code"], tr(c["label"]), num(c["expected"]),
                               num(c["observed"]), num(c["delta"]), tr(c["status"]), tr(c["note"]), None, None])
        for i in a["issues"]:
            append(queue, [did, i["row_id"], i["page"], i["code"], tr(i["message"]), None, None, None, tr("待复核"), None, None, None])
        for i in a.get("resolved_issues", []):
            append(queue, [did, i["row_id"], i["page"], i["code"], tr(i["message"]), None, None, None, tr("人工已解决"),
                           i["resolution_note"], i["reviewer"], i["reviewed_at"]])
        for key, value in f.items():
            evidence = d["field_evidence"].get(key, [])
            append(fields, [did, key, num(value) if key in AMOUNTS else value, ", ".join(str(e["page"]) for e in evidence),
                            " / ".join(e["text"] for e in evidence) or tr("见原 PDF 和 source.json；推断字段/模型或人工提取无独立字段定位")])
        append(fields, [did, "source_sha256", d["source"]["sha256"], None, tr("引擎：") + d["engine"]])
        append(fields, [did, "calculation_policy", tr(f"容差={a['tolerance']}；金额小数位={a['decimal_places']}"), None, a["rounding"]])
        if d.get("human_review"):
            r = d["human_review"]
            append(fields, [did, "human_review", r["reviewer"], None,
                            tr(f"{r['reviewed_at']}；完整性确认={r['coverage_confirmed']}；{r['notes']}")])
    for e in errors:
        append(overview, [e["document_id"], e["file"], None, None, None, None, None, None, None, tr("待人工处理"), tr("处理失败")])
        append(queue, [e["document_id"], None, None, "PROCESSING_ERROR", tr(e["message"]), None, None, None, tr("无法核验"), e["file"], None, None])
    finish(overview, integer_cols=(6, 7, 8, 9))
    finish(detail, (6, 7, 8, 9), integer_cols=(3,), general_cols=(5,))
    finish(queue, (6, 7, 8), (11, 12))
    finish(fields)
    for row in fields.iter_rows(min_row=4):
        if row[1].value in AMOUNTS:
            row[2].number_format = '#,##0.00######;[Red](#,##0.00######);0.00'
    wb.properties.creator = "SheetAudit AI"
    wb.properties.description = "Local audit snapshot. Human review required."
    wb.save(output)
