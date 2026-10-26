import io
import json
import tempfile
import unittest
from copy import deepcopy
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import load_workbook
from sheetaudit.audit import audit_document
from sheetaudit.cli import main
from sheetaudit.extract import extract_document, extract_rules
from sheetaudit.model import FIELDS, identified_issues
from sheetaudit.privacy import private_summary
from sheetaudit.review import apply_review, review_template
from sheetaudit.review_excel import read_review_workbook, write_review_workbook

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "examples/pdfs/good_invoice.pdf"
REG = ROOT / "examples/regressions"
HEADER = ["Description", "Quantity", "Unit Price", "Amount"]


def document():
    doc, _ = extract_document(PDF)
    doc["document_id"] = "d001_" + doc["source"]["sha256"][:8]
    return doc


def signed(doc, resolved=False):
    edit = review_template(doc)
    edit.update(reviewer="Synthetic reviewer", reviewed_at="2026-10-25T10:00:00+08:00", notes="Synthetic test only", coverage_confirmed=True)
    if resolved:
        for entry in edit["resolutions"]:
            entry.update(action="resolved", note="Confirmed against synthetic source")
    return edit


def fill_workbook(path):
    wb = load_workbook(path)
    ws = wb["复核签注"]
    ws["B4"] = "Synthetic reviewer"
    ws["B5"] = "2026-10-25T10:00:00+08:00"
    ws["B6"] = "已确认"
    ws["B7"] = "Synthetic test only"
    for row in wb["问题处理"].iter_rows(min_row=4):
        row[5].value = "已解决"
        row[6].value = "Checked synthetic source"
    return wb


class LayoutRegressionTests(unittest.TestCase):
    def test_product_prefixes_and_exact_total_are_not_summaries(self):
        names = ["Shipping labels", "Shipping: labels", "Tax stamp", "Total protection", "Discount voucher", "Total"]
        rows = [[name, "2", "10", "20"] for name in names]
        content, warnings, _ = extract_rules([{"page": 1, "text": "", "tables": [[HEADER, *rows]]}], "en")
        self.assertEqual([r["description"] for r in content["items"]], names)

    def test_summary_in_middle_does_not_erase_mapping(self):
        rows = [HEADER, ["A", "1", "2", "2"], ["Subtotal", "", "", "2"], ["B", "1", "3", "3"]]
        content, _, _ = extract_rules([{"page": 1, "text": "", "tables": [rows]}], "en")
        self.assertEqual(len(content["items"]), 2)
        self.assertEqual(content["fields"]["subtotal"], "2")

    def test_actual_pdf_regressions(self):
        for filename, rows in [("shipping_product.pdf", 3), ("continuation_grid.pdf", 2),
                               ("borderless.pdf", 2), ("continuation_borderless.pdf", 2)]:
            with self.subTest(filename=filename):
                doc, _ = extract_document(REG / filename)
                self.assertEqual(len(doc["items"]), rows)
                self.assertEqual(audit_document(doc)["differences"], 0)
                self.assertEqual(audit_document(doc)["uncheckable"], 0)
                if filename.startswith("continuation"):
                    self.assertEqual([i["page"] for i in doc["items"]], [1, 2])
                    self.assertIn("INHERITED_HEADER", {i["code"] for i in doc["extraction_issues"]})

    def test_changed_column_positions_not_inherited(self):
        pages = [{"page": 1, "text": "", "tables": [[HEADER, ["A", "1", "2", "2"]]], "table_edges": [[0, 100, 200, 300, 400]]},
                 {"page": 2, "text": "", "tables": [[["B", "1", "3", "3"]]], "table_edges": [[0, 50, 100, 150, 200]]}]
        content, warnings, _ = extract_rules(pages, "en")
        self.assertEqual(len(content["items"]), 1)
        self.assertIn("NO_ITEMS_ON_PAGE", {i["code"] for i in warnings})

    def test_new_document_number_blocks_inheritance(self):
        pages = [{"page": 1, "text": "Invoice No: A", "tables": [[HEADER, ["A", "1", "2", "2"]]]},
                 {"page": 2, "text": "Invoice No: B", "tables": [[["B", "1", "3", "3"]]]}]
        content, warnings, _ = extract_rules(pages, "en")
        self.assertEqual(len(content["items"]), 1)
        self.assertIn("CONFLICTING_FIELD", {i["code"] for i in warnings})


class ResolutionTests(unittest.TestCase):
    def test_resolved_notices_leave_open_queue_but_keep_history(self):
        doc = document()
        updated = apply_review(doc, signed(doc, True), PDF)
        result = audit_document(updated)
        self.assertEqual(result["issues"], [])
        self.assertEqual(len(result["resolved_issues"]), 1)
        self.assertIn("本次复核完成", result["review_status"])
        self.assertTrue(updated["extraction_issues"])

    def test_data_change_reopens_previous_decisions(self):
        doc = document()
        updated = apply_review(doc, signed(doc, True), PDF)
        updated["items"][0]["amount"] = "39.00"
        result = audit_document(updated)
        self.assertEqual(result["resolved_issues"], [])
        self.assertTrue(result["issues"])
        self.assertGreater(result["differences"], 0)
        self.assertIn("需重新复核", result["review_status"])

    def test_arithmetic_difference_cannot_be_manually_dismissed(self):
        doc = document()
        edit = signed(doc, True)
        edit["items"][0]["amount"] = "40"
        result = audit_document(apply_review(doc, edit, PDF))
        self.assertEqual(len(result["issues"]), 0)
        self.assertGreater(result["differences"], 0)
        self.assertIn("待处理", result["review_status"])

    def test_unknown_resolution_and_missing_reason_rejected(self):
        doc = document()
        for mutation in (lambda e: e["resolutions"][0].update(issue_id="made-up"),
                         lambda e: e["resolutions"][0].update(note=""),
                         lambda e: e.update(coverage_confirmed=False)):
            edit = signed(doc, True)
            mutation(edit)
            with self.assertRaises(ValueError):
                apply_review(doc, edit, PDF)

    def test_partial_review_does_not_claim_completeness(self):
        doc = document()
        edit = signed(doc)
        edit["coverage_confirmed"] = False
        result = audit_document(apply_review(doc, edit, PDF))
        self.assertIn("完整性未确认", result["review_status"])


class ReviewExcelTests(unittest.TestCase):
    def test_roundtrip_numbers_signoff_and_resolution(self):
        doc = document()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path, PDF)
            wb = fill_workbook(path)
            wb["明细修正"]["F4"] = 42.50
            wb.save(path)
            edit = read_review_workbook(path, doc)
            self.assertEqual(edit["items"][0]["amount"], "42.5")
            result = audit_document(apply_review(doc, edit, PDF))
            self.assertEqual(result["differences"], 2)
            self.assertEqual(len(result["resolved_issues"]), 1)

    def test_formulas_rejected_even_with_cached_or_blank_value(self):
        doc = document()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path)
            wb = fill_workbook(path)
            wb["明细修正"]["D4"] = "=1+2"
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "公式"):
                read_review_workbook(path, doc)

    def test_stale_workbook_and_original_column_change_rejected(self):
        doc = document()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path)
            stale = deepcopy(doc)
            stale["fields"]["buyer"] = "changed"
            with self.assertRaises(ValueError):
                read_review_workbook(path, stale)
            wb = fill_workbook(path)
            wb["字段修正"]["C4"] = "modified original"
            wb.save(path)
            with self.assertRaises(ValueError):
                read_review_workbook(path, doc)

    def test_column_reordering_rejected(self):
        doc = document()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path)
            wb = fill_workbook(path)
            wb["明细修正"]["D3"], wb["明细修正"]["E3"] = "单价", "数量"
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "表头"):
                read_review_workbook(path, doc)

    def test_delete_is_explicit_and_addition_has_evidence(self):
        doc = document()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path)
            wb = fill_workbook(path)
            wb["明细修正"]["H4"] = "删除"
            wb["明细修正"].append(["new1", 1, "Replacement", "3", "12.50", "37.50", "Synthetic evidence", "保留"])
            wb.save(path)
            edit = read_review_workbook(path, doc)
            self.assertEqual(len(edit["items"]), 2)
            self.assertEqual(edit["items"][-1]["row_id"], "new1")
            apply_review(doc, edit, PDF)
            wb["明细修正"].delete_rows(4)
            wb.save(path)
            with self.assertRaises(ValueError):
                read_review_workbook(path, doc)

    def test_empty_document_can_be_manually_entered(self):
        scanned = ROOT / "examples/pdfs/scanned_invoice.pdf"
        doc, _ = extract_document(scanned)
        doc["document_id"] = "d001_" + doc["source"]["sha256"][:8]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            write_review_workbook(doc, path)
            wb = fill_workbook(path)
            wb["明细修正"].append(["manual1", 1, "Cable", "2", "10.00", "20.00", "Cable | 2 | 10.00 | 20.00", "保留"])
            wb.save(path)
            edit = read_review_workbook(path, doc)
            self.assertEqual(len(edit["items"]), 1)
            apply_review(doc, edit, scanned)

    def test_cli_excel_review_and_policy_preservation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.assertEqual(main(["audit", str(PDF), "--out", str(root / "run"), "--tolerance", "0"]), 0)
            folder = next((root / "run/documents").iterdir())
            path = folder / "review.xlsx"
            fill_workbook(path).save(path)
            self.assertEqual(main(["review", str(folder / "extracted.json"), "--corrections", str(path),
                                   "--source", str(PDF), "--out", str(root / "reviewed"), "--fail-on-issues"]), 0)
            result = json.loads((root / "reviewed/summary.json").read_text(encoding="utf-8"))
            self.assertEqual(result["tolerance"], "0")
            self.assertEqual(result["review_notices"], 0)
            self.assertEqual(result["resolved_notices"], 1)


class PrivacyTests(unittest.TestCase):
    def test_allowlist_drops_all_free_text_and_unknown_issue_codes(self):
        d = document()
        secret = "PRIVATE_CUSTOMER_987654321"
        d["fields"]["buyer"] = secret
        d["source"]["file"] = secret + ".pdf"
        d["extraction_issues"].append({"code": secret, "message": secret, "page": 1, "row_id": secret})
        summary = {"records": [{"document": d, "audit": audit_document(d)}], "errors": [{"file": secret, "message": secret}]}
        safe = private_summary(summary)
        encoded = json.dumps(safe)
        self.assertNotIn(secret, encoded)
        self.assertNotIn(d["source"]["sha256"], encoded)
        self.assertNotIn("37.50", encoded)
        self.assertEqual(safe["issue_counts"]["OTHER"], 1)
        self.assertEqual(safe["failed"], 1)

    def test_summary_cli_exposes_statistics_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out"
            main(["audit", str(PDF), "--out", str(out)])
            captured = io.StringIO()
            with redirect_stdout(captured):
                self.assertEqual(main(["summary", str(out)]), 0)
            value = json.loads(captured.getvalue())
            self.assertEqual(value["items"], 2)
            self.assertNotIn("INV-DEMO", captured.getvalue())
            self.assertNotIn("good_invoice.pdf", captured.getvalue())


if __name__ == "__main__":
    unittest.main()
