import json
import tempfile
import unittest
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch, MagicMock

from openpyxl import load_workbook
from sheetaudit.audit import audit_document
from sheetaudit.cli import main
from sheetaudit.extract import extract_document, extract_rules, extract_ollama
from sheetaudit.model import FIELDS, parse_number, validate_content
from sheetaudit.review import review_template, apply_review

ROOT = Path(__file__).resolve().parents[1]
PDFS = ROOT / "examples" / "pdfs"


def sample():
    doc = extract_document(PDFS / "good_invoice.pdf")[0]
    doc["document_id"] = "d001_" + doc["source"]["sha256"][:8]
    return doc


class NumberTests(unittest.TestCase):
    def test_locales_and_credits(self):
        self.assertEqual(parse_number("$1,234.50"), "1234.50")
        self.assertEqual(parse_number("1.234,50", "de"), "1234.50")
        self.assertEqual(parse_number("(12.50)"), "-12.50")
        self.assertEqual(parse_number("０.５０"), "0.50")
        self.assertEqual(parse_number("0"), "0")
        self.assertIsNone(parse_number(""))

    def test_ambiguous_and_unsafe_numbers(self):
        for value in ("1,23", "O2", "2 pcs", "NaN", "Infinity", "1e3", "1.234,50", "1/2", "=1+2", "9"*16):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_number(value)


class ExtractionTests(unittest.TestCase):
    def test_clean_pdf(self):
        doc = sample()
        self.assertEqual(doc["fields"]["document_number"], "INV-DEMO-001")
        self.assertEqual(doc["fields"]["total"], "63.25")
        self.assertEqual(len(doc["items"]), 2)
        result = audit_document(doc)
        self.assertEqual(result["differences"], 0)
        self.assertEqual(result["uncheckable"], 0)
        self.assertEqual(result["review_status"], "待人工复核")

    def test_chinese_difference(self):
        doc, _ = extract_document(ROOT / "tests/fixtures/discrepancy_invoice_zh.pdf")
        self.assertEqual(doc["fields"]["seller"], "示例办公用品公司")
        self.assertEqual(doc["fields"]["currency"], "CNY")
        result = audit_document(doc)
        self.assertEqual(result["differences"], 1)
        self.assertEqual(Decimal(result["checks"][0]["delta"]), Decimal("5"))

    def test_repeated_lines_across_pages_preserved(self):
        doc, _ = extract_document(PDFS / "purchase_order_multipage.pdf")
        self.assertEqual([i["page"] for i in doc["items"]], [1, 2])
        self.assertEqual(doc["items"][0]["description"], doc["items"][1]["description"])
        self.assertEqual(audit_document(doc)["differences"], 0)

    def test_scanned_pdf_cannot_pass_as_empty_clean(self):
        doc, _ = extract_document(PDFS / "scanned_invoice.pdf")
        result = audit_document(doc)
        self.assertEqual(len(doc["items"]), 0)
        self.assertGreater(result["uncheckable"], 0)
        self.assertIn("NO_TEXT", {i["code"] for i in result["issues"]})

    def test_missing_values_not_repaired_from_amount(self):
        doc, _ = extract_document(PDFS / "unsupported_invoice.pdf")
        self.assertEqual(len(doc["items"]), 2)
        self.assertIsNone(doc["items"][0]["quantity"])
        self.assertIsNone(doc["items"][1]["unit_price"])
        result = audit_document(doc)
        self.assertEqual(result["uncheckable"], 3)

    def test_pipe_fallback_and_no_cross_strategy_dedup(self):
        text = "Invoice No: TEST\nDescription | Qty | Unit Price | Amount\nCable | 2 | 10 | 20\nCable | 2 | 10 | 20\nSubtotal: 40"
        content, issues, _ = extract_rules([{"page": 1, "text": text, "tables": []}], "en")
        self.assertEqual(len(content["items"]), 2)
        self.assertEqual(content["fields"]["subtotal"], "40")

    def test_default_does_not_open_network(self):
        with patch("socket.socket", side_effect=AssertionError("Network forbidden")):
            self.assertEqual(len(sample()["items"]), 2)

    def test_ocr_route_with_mock_engine(self):
        text = "Invoice No: SCAN\nDescription | Qty | Unit Price | Amount\nCable | 2 | 10.00 | 20.00\nSubtotal: 20.00"
        with patch("sheetaudit.extract.ocr_page", return_value=text) as ocr:
            doc, pages = extract_document(PDFS / "scanned_invoice.pdf", ocr=True)
        self.assertEqual(len(doc["items"]), 1)
        self.assertEqual(pages[0]["method"], "ocr")
        self.assertIn("OCR_REVIEW", {i["code"] for i in doc["extraction_issues"]})
        ocr.assert_called_once()

    def test_ocr_dependency_failure(self):
        from sheetaudit.extract import ocr_page
        with patch("sheetaudit.extract.shutil.which", return_value=None), self.assertRaises(ValueError):
            ocr_page(PDFS / "scanned_invoice.pdf", 1, "eng", "missing-tesseract")


class AuditTests(unittest.TestCase):
    def test_decimal_rounding_and_tolerance_boundaries(self):
        doc = sample()
        doc["items"] = [dict(doc["items"][0], quantity="1", unit_price="1.005", amount="1.02")]
        c = audit_document(doc)["checks"][0]
        self.assertEqual(c["expected"], "1.01")
        self.assertEqual(c["status"], "算术一致")
        self.assertEqual(audit_document(doc, "0.009")["checks"][0]["status"], "差异")
        doc["items"][0].update(quantity="-1", unit_price="1.005", amount="-1.01")
        self.assertEqual(audit_document(doc, "0")["checks"][0]["status"], "算术一致")

    def test_zero_and_missing_are_different(self):
        doc = sample()
        doc["items"][0]["quantity"] = "0"
        self.assertEqual(audit_document(doc)["checks"][0]["status"], "差异")
        doc["items"][0]["quantity"] = None
        self.assertEqual(audit_document(doc)["checks"][0]["status"], "无法核验")
        doc["items"][0]["amount"] = None
        self.assertEqual(audit_document(doc)["checks"][-2]["status"], "无法核验")

    def test_total_adjustments_and_subtotal_independent(self):
        doc = sample()
        doc["fields"].update(subtotal="59", tax="6", discount="2", shipping="3", total="66")
        checks = audit_document(doc)["checks"]
        self.assertEqual(checks[-2]["status"], "差异")
        self.assertEqual(checks[-1]["status"], "算术一致")
        doc["fields"]["shipping"] = None
        self.assertEqual(audit_document(doc)["checks"][-1]["status"], "无法核验")

    def test_currency_precision_and_missing(self):
        doc = sample()
        doc["fields"]["currency"] = "JPY"
        self.assertIn("CURRENCY_PRECISION", {i["code"] for i in audit_document(doc)["issues"]})
        self.assertNotIn("CURRENCY_PRECISION", {i["code"] for i in audit_document(doc, places=0)["issues"]})

    def test_invalid_tolerance(self):
        for tolerance in ("NaN", "-1", "Infinity", "1001"):
            with self.subTest(tolerance=tolerance), self.assertRaises(ValueError):
                audit_document(sample(), tolerance)


class ReviewTests(unittest.TestCase):
    def correction(self):
        doc = sample()
        correction = review_template(doc)
        correction.update(reviewer="Demo reviewer", reviewed_at="2026-10-24T20:00:00+08:00",
                          notes="Verified with synthetic source", coverage_confirmed=True)
        return doc, correction

    def test_review_preserves_before_and_recalculates(self):
        doc, correction = self.correction()
        correction["items"][0]["amount"] = "38.50"
        updated = apply_review(doc, correction, PDFS / "good_invoice.pdf")
        self.assertEqual(doc["items"][0]["amount"], "37.50")
        self.assertEqual(updated["human_review"]["changes"][0]["section"], "items")
        self.assertEqual(audit_document(updated)["differences"], 2)
        self.assertIn("非自动批准", audit_document(updated)["review_status"])

    def test_wrong_source_stale_base_and_missing_signoff_rejected(self):
        doc, correction = self.correction()
        with self.assertRaises(ValueError):
            apply_review(doc, correction, PDFS / "scanned_invoice.pdf")
        doc["fields"]["buyer"] = "Changed"
        with self.assertRaises(ValueError):
            apply_review(doc, correction, PDFS / "good_invoice.pdf")
        doc, correction = self.correction()
        correction["reviewer"] = ""
        with self.assertRaises(ValueError):
            apply_review(doc, correction, PDFS / "good_invoice.pdf")

    def test_schema_rejects_numeric_float_unknown_fields_invalid_page(self):
        doc = sample()
        content = {"fields": doc["fields"], "items": doc["items"]}
        for key, value in (("quantity", 1.1), ("page", 999), ("row_id", "../escape")):
            bad = deepcopy(content)
            bad["items"][0][key] = value
            with self.subTest(key=key), self.assertRaises(ValueError):
                validate_content(bad, 1)

    def test_review_document_id_cannot_escape_output_directory(self):
        doc, correction = self.correction()
        doc["document_id"] = "../../outside"
        with self.assertRaises(ValueError):
            apply_review(doc, correction, PDFS / "good_invoice.pdf")


class CliTests(unittest.TestCase):
    def test_batch_report_contents_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "report"
            self.assertEqual(main(["audit", str(PDFS), "--out", str(out), "--language", "zh"]), 0)
            summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual(summary["documents"], 5)
            self.assertEqual(summary["failed"], 0)
            self.assertEqual(summary["differences"], 1)
            wb = load_workbook(out / "report.xlsx")
            self.assertEqual(wb.sheetnames, ["单据汇总", "明细核对", "差异与复核", "字段与来源"])
            self.assertEqual(wb["明细核对"].max_row, 11)
            before = (out / "summary.json").read_bytes()
            self.assertEqual(main(["audit", str(PDFS), "--out", str(out)]), 2)
            self.assertEqual(before, (out / "summary.json").read_bytes())

    def test_failed_pdf_does_not_hide_other_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "bad.pdf"
            bad.write_bytes(b"not a pdf")
            out = Path(tmp) / "result"
            self.assertEqual(main(["audit", str(bad), str(PDFS / "good_invoice.pdf"), "--out", str(out)]), 2)
            summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual((summary["documents"], summary["failed"]), (1, 1))

    def test_strict_exit_requires_review_even_if_arithmetic_matches(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main(["audit", str(PDFS / "good_invoice.pdf"), "--out", str(Path(tmp) / "out"), "--fail-on-issues"]), 1)

    def test_excel_injection_is_text(self):
        from sheetaudit.report import write_report
        doc = sample()
        doc["document_id"] = "d001"
        doc["items"][0]["description"] = '=HYPERLINK("https://example.com", "click")\x01'
        doc["fields"]["seller"] = "=1+2"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            write_report([{"document": doc, "audit": audit_document(doc)}], [], path)
            wb = load_workbook(path)
            self.assertEqual(wb["明细核对"]["D4"].data_type, "s")
            self.assertFalse(any(c.data_type == "f" for ws in wb for row in ws for c in row))

    def test_duplicate_files_are_flagged_but_kept(self):
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / "renamed.pdf"
            shutil.copyfile(PDFS / "good_invoice.pdf", copy)
            out = Path(tmp) / "result"
            self.assertEqual(main(["audit", str(copy), str(PDFS / "good_invoice.pdf"), "--out", str(out)]), 0)
            summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual(summary["documents"], 2)
            for record in summary["records"]:
                self.assertIn("DUPLICATE_FILE", {i["code"] for i in record["audit"]["issues"]})

    def test_empty_directory_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main(["audit", tmp, "--out", str(Path(tmp) / "out")]), 2)

    def test_review_cli_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.assertEqual(main(["audit", str(PDFS / "good_invoice.pdf"), "--out", str(root / "first")]), 0)
            folder = next((root / "first" / "documents").iterdir())
            correction = json.loads((folder / "review-template.json").read_text(encoding="utf-8"))
            correction.update(reviewer="Demo", reviewed_at="2026-10-24T20:00:00+08:00", notes="Source verified", coverage_confirmed=True)
            edit = root / "correction.json"
            edit.write_text(json.dumps(correction), encoding="utf-8")
            self.assertEqual(main(["review", str(folder / "extracted.json"), "--source", str(PDFS / "good_invoice.pdf"),
                                   "--corrections", str(edit), "--out", str(root / "second")]), 0)
            self.assertTrue(next((root / "second" / "documents").iterdir()).joinpath("before-review.json").exists())


class OllamaTests(unittest.TestCase):
    def test_contract_with_mock_local_response(self):
        doc = sample()
        content = {"fields": doc["fields"], "items": doc["items"]}
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({"done": True, "response": json.dumps(content)}).encode()
        with patch("sheetaudit.extract.build_opener") as build:
            build.return_value.open.return_value = response
            result = extract_ollama([{"page": 1, "text": "demo", "tables": []}], "local-demo", 11434)
            self.assertEqual(len(result["items"]), 2)
            request = build.return_value.open.call_args.args[0]
            self.assertEqual(request.full_url, "http://127.0.0.1:11434/api/generate")

    def test_cloud_name_and_truncation_rejected(self):
        with self.assertRaises(ValueError):
            extract_ollama([], "some-cloud", 11434)
        with self.assertRaises(ValueError):
            extract_ollama([{"text": "x" * 45001}], "local-demo", 11434)

    def test_incomplete_and_invalid_model_output_rejected(self):
        for envelope in ({"done": False, "response": "{}"},
                         {"done": True, "done_reason": "length", "response": "{}"},
                         {"done": True, "response": '{"fields": {}, "items": []}'}):
            response = MagicMock()
            response.__enter__.return_value.read.return_value = json.dumps(envelope).encode()
            with patch("sheetaudit.extract.build_opener") as build, self.assertRaises(ValueError):
                build.return_value.open.return_value = response
                extract_ollama([{"page": 1, "text": "demo", "tables": []}], "local-demo", 11434)


if __name__ == "__main__":
    unittest.main()
