import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from sheetaudit.desktop import audit_arguments, execute_job, new_output, review_arguments, default_output_directory

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "examples/pdfs/good_invoice.pdf"


class DesktopWorkflowTests(unittest.TestCase):
    def test_portable_output_follows_program_folder_when_moved(self):
        import sys
        with tempfile.TemporaryDirectory(prefix="桌面 中文 ") as folder:
            base = Path(folder).resolve()
            with patch.object(sys, "frozen", True, create=True):
                with patch.object(sys, "executable", str(base / "程序" / "SheetAudit.exe")):
                    self.assertEqual(default_output_directory(), base / "输出报告")
                with patch.object(sys, "executable", str(base / "其他位置" / "SheetAudit.exe")):
                    self.assertEqual(default_output_directory(), base / "其他位置" / "输出报告")

    def test_invalid_settings_fail_before_creating_output(self):
        with tempfile.TemporaryDirectory() as folder:
            for value in ("", "NaN", "Infinity", "-0.01", "1001", "abc"):
                with self.subTest(value=value), self.assertRaises(ValueError):
                    audit_arguments([str(PDF)], folder, value)
            with self.assertRaises(ValueError):
                audit_arguments([str(PDF)], folder, places="9")
            with self.assertRaises(ValueError):
                audit_arguments([], folder)
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_unique_paths_and_local_job_with_progress(self):
        with tempfile.TemporaryDirectory(prefix="中文 空格 ") as folder:
            self.assertNotEqual(new_output(folder), new_output(folder))
            events = []
            args = audit_arguments([str(PDF)], folder, "0", "3")
            with patch("socket.socket", side_effect=AssertionError("No network")):
                result = execute_job(args, lambda *event: events.append(event))
            self.assertEqual(result["code"], 0)
            self.assertEqual(result["summary"]["documents"], 1)
            self.assertEqual(result["summary"]["items"], 2)
            self.assertGreater(len(events), 1)
            self.assertTrue((Path(result["output"]) / "report.xlsx").is_file())
            self.assertNotIn("INV-DEMO", str(result["summary"]))

    def test_partial_failure_preserves_success_and_error_count(self):
        with tempfile.TemporaryDirectory() as folder:
            bad = Path(folder) / "broken.pdf"
            bad.write_bytes(b"not a PDF")
            result = execute_job(audit_arguments([str(PDF), str(bad)], folder))
            self.assertEqual(result["code"], 2)
            self.assertEqual(result["summary"]["documents"], 1)
            self.assertEqual(result["summary"]["failed"], 1)

    def test_review_signoff_and_source_fingerprint_still_required(self):
        with tempfile.TemporaryDirectory() as folder:
            result = execute_job(audit_arguments([str(PDF)], folder, "0", "3"))
            extracted = next(Path(result["output"]).glob("documents/*/extracted.json"))
            workbook = extracted.with_name("review.xlsx")
            args = review_arguments(str(extracted), str(workbook), str(PDF), folder)
            with self.assertRaises(ValueError):
                execute_job(args)
            wb = load_workbook(workbook)
            ws = wb["复核签注"]
            ws["B4"], ws["B5"] = "Synthetic test reviewer", "2026-10-25T20:00:00+08:00"
            ws["B6"], ws["B7"] = "已确认", "Synthetic regression only"
            for row in wb["问题处理"].iter_rows(min_row=4):
                row[5].value, row[6].value = "已解决", "Checked synthetic source"
            wb.save(workbook)
            reviewed = execute_job(args)
            # Three decimals on a USD invoice still needs currency-policy review.
            self.assertEqual(reviewed["summary"]["issue_counts"], {"CURRENCY_PRECISION": 1})
            self.assertEqual(reviewed["summary"]["resolved_notices"], 1)
            import json
            record = json.loads(next(Path(reviewed["output"]).glob("documents/*/extracted.json")).read_text(encoding="utf-8"))
            self.assertEqual(record["audit_policy"], {"tolerance": "0", "decimal_places": 3})
            wrong = ROOT / "examples/pdfs/discrepancy_invoice_zh.pdf"
            with self.assertRaises(ValueError):
                execute_job(review_arguments(str(extracted), str(workbook), str(wrong), folder))


if __name__ == "__main__":
    unittest.main()
