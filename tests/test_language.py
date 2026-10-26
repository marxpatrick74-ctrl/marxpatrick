"""Language selection must preserve arithmetic, source data and review integrity."""
import contextlib
import io
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from sheetaudit.audit import audit_document
from sheetaudit.cli import main
from sheetaudit.desktop import audit_arguments, execute_job
from sheetaudit.extract import extract_document
from sheetaudit.i18n import translate, use_language, current_language, load_language, save_language
from sheetaudit.model import json_hash
from sheetaudit.report import write_report
from sheetaudit.review import apply_review
from sheetaudit.review_excel import write_review_workbook, read_review_workbook

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "examples/pdfs/good_invoice.pdf"


def document():
    doc = extract_document(PDF)[0]
    doc["document_id"] = "d001_" + doc["source"]["sha256"][:8]
    return doc


def sign(path, language):
    book = load_workbook(path)
    name = lambda key: translate(key, language)
    sheet = book[name("复核签注")]
    sheet["B4"] = "Synthetic reviewer"
    sheet["B5"] = "2026-10-26T12:00:00+08:00"
    sheet["B6"] = name("已确认")
    sheet["B7"] = "Checked all synthetic pages"
    for row in book[name("问题处理")].iter_rows(min_row=4):
        row[5].value, row[6].value = name("已解决"), "Checked synthetic source"
    book.save(path)
    book.close()


class LanguageWorkflowTests(unittest.TestCase):
    def test_english_default_cli_and_all_public_demo_workbooks(self):
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()) as output:
            out = Path(folder) / "en"
            self.assertEqual(main(["audit", str(ROOT / "examples/pdfs"), "--out", str(out)]), 0)
            self.assertIn("Report generated", output.getvalue())
            paths = [out / "report.xlsx", *out.glob("documents/*/review.xlsx")]
            for path in paths:
                book = load_workbook(path)
                self.assertFalse(re.search(r"[\u4e00-\u9fff]", " ".join(book.sheetnames)))
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value, str):
                                self.assertFalse(re.search(r"[\u4e00-\u9fff]", cell.value), (path, sheet.title, cell.coordinate, cell.value))
                book.close()
            summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual((summary["language"], summary["differences"], summary["failed"]), ("en", 1, 0))

    def test_language_before_and_after_command_and_help(self):
        for arguments, expected in ((["--language", "en", "audit", "--help"], "Extract PDFs"),
                                    (["audit", "--language=zh", "--help"], "提取 PDF")):
            with self.subTest(arguments=arguments), contextlib.redirect_stdout(io.StringIO()) as output:
                with self.assertRaises(SystemExit) as exit:
                    main(arguments)
                self.assertEqual(exit.exception.code, 0)
                self.assertIn(expected, output.getvalue())

    def test_errors_follow_cli_language(self):
        for language, expected in (("en", "Inputs must"), ("zh", "输入必须")):
            with self.subTest(language=language), contextlib.redirect_stderr(io.StringIO()) as output:
                self.assertEqual(main(["audit", "missing.pdf", "--out", "unused", "--language", language]), 2)
                self.assertIn(expected, output.getvalue())

    def test_both_workbook_languages_can_be_read_from_other_ui_language(self):
        for language in ("en", "zh"):
            with self.subTest(language=language), tempfile.TemporaryDirectory() as folder:
                doc = document()
                original_hash = json_hash(doc)
                path = Path(folder) / "review.xlsx"
                with use_language(language):
                    write_review_workbook(doc, path)
                sign(path, language)
                with use_language("zh" if language == "en" else "en"):
                    correction = read_review_workbook(path, doc)
                    reviewed = apply_review(doc, correction, PDF)
                self.assertEqual(json_hash(doc), original_hash)
                self.assertEqual(audit_document(reviewed)["issues"], [])
                self.assertEqual(audit_document(reviewed)["differences"], 0)

    def test_legacy_chinese_workbook_without_language_metadata(self):
        with tempfile.TemporaryDirectory() as folder:
            doc, path = document(), Path(folder) / "legacy.xlsx"
            write_review_workbook(doc, path)
            book = load_workbook(path)
            meta = book["_meta"]
            for row in list(meta):
                if row[0].value == "language":
                    meta.delete_rows(row[0].row)
            book.save(path)
            book.close()
            sign(path, "zh")
            with use_language("en"):
                self.assertTrue(read_review_workbook(path, doc)["coverage_confirmed"])

    def test_localization_does_not_translate_source_values_or_evidence(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = document()
            doc["fields"]["seller"] = "差异"  # Deliberately matches an application label.
            doc["items"][0]["description"] = "已确认"
            doc["items"][0]["evidence"] = "数量×单价"
            before = json_hash(doc)
            report, review = Path(folder) / "report.xlsx", Path(folder) / "review.xlsx"
            with use_language("en"):
                write_report([{"document": doc, "audit": audit_document(doc)}], [], report)
                write_review_workbook(doc, review)
            book = load_workbook(report)
            self.assertEqual(book["Line checks"]["D4"].value, "已确认")
            self.assertEqual(book["Line checks"]["L4"].value, "数量×单价")
            self.assertIn("差异", [row[2].value for row in book["Fields and sources"].iter_rows(min_row=4)])
            book.close()
            sign(review, "en")
            correction = read_review_workbook(review, doc)
            self.assertEqual(correction["fields"]["seller"], "差异")
            self.assertEqual(correction["items"][0]["description"], "已确认")
            self.assertEqual(json_hash(doc), before)

    def test_english_review_rejects_formulas_and_modified_original_notices(self):
        with tempfile.TemporaryDirectory() as folder:
            doc, path = document(), Path(folder) / "review.xlsx"
            for target in ("formula", "notice", "language"):
                with use_language("en"):
                    write_review_workbook(doc, path)
                sign(path, "en")
                book = load_workbook(path)
                if target == "formula": book["Field corrections"]["D4"] = "=1+1"
                if target == "notice": book["Issue resolutions"]["E4"] = "Altered original notice"
                if target == "language": book["_meta"]["B2"] = "fr"
                book.save(path)
                book.close()
                with self.subTest(target=target), self.assertRaises(ValueError):
                    read_review_workbook(path, doc)

    def test_desktop_report_language_comes_from_job_arguments(self):
        with tempfile.TemporaryDirectory() as folder:
            args = audit_arguments([str(PDF)], folder, language="en")
            with use_language("zh"):
                result = execute_job(args)
            book = load_workbook(Path(result["output"]) / "report.xlsx")
            self.assertIn("Documents", book.sheetnames)
            book.close()

    def test_language_preference_is_local_only_and_survives_invalid_files(self):
        with tempfile.TemporaryDirectory(prefix="语言 ") as folder:
            path = Path(folder) / "settings.json"
            self.assertEqual(load_language(path), "en")
            self.assertTrue(save_language("zh", path))
            self.assertEqual(load_language(path), "zh")
            self.assertEqual(json.loads(path.read_text()), {"language": "zh"})
            for data in ('[]', '{"language":"fr"}', 'not JSON'):
                path.write_text(data)
                self.assertEqual(load_language(path), "en")
            with patch("pathlib.Path.mkdir", side_effect=PermissionError):
                self.assertFalse(save_language("en", path))

    def test_context_language_is_restored_after_error(self):
        self.assertEqual(current_language(), "zh")
        with self.assertRaises(RuntimeError):
            with use_language("en"):
                self.assertEqual(translate("差异"), "Difference")
                raise RuntimeError()
        self.assertEqual(current_language(), "zh")


if __name__ == "__main__":
    unittest.main()
